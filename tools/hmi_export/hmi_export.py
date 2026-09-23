# -*- coding: utf-8 -*-
"""hmi_export.py —— QT上位机地图 → 大彩组态屏CSV 导出工具 (GUI + CLI)。

输入(二选一):
  A. QT上位机 SQLite 数据库 data.db (Map/MapBlock/RowAttributes/MapAisle/ABLine表)
  B. QT上位机导出的 地图<名>.xlsx (信息/AB线/编号/密度/是否播种 五工作表,
     布局对齐 Map::exportToXlsx, common/map.cpp:181)

输出: 组态屏可识别的 map_<slug>.csv, 格式与 HIMproject csv.lua render_map 一致:
  #MAP,<显示名(可中文)> / #DIM,rows,cols,machine_cols,col_gap[,block_col]
  #ROWLEN,... / #AISLE,...(rows+1条) / [#AB,...] / #BLOCKS r,c,num,density,sow(1基)
  编码GB2312 + CRLF; 文件名纯ASCII(屏端文件系统限制), slug算法与 store.lua 一致。

用法:
  python hmi_export.py                      # 图形界面
  python hmi_export.py --db data.db --out D:\\usb
  python hmi_export.py --db data.db --out . --map 对照田
  python hmi_export.py --xlsx 地图对照田.xlsx 地图乙.xlsx --out .
  python hmi_export.py --selftest           # 自测(含屏端csv.lua回读验证)
"""
import argparse
import math
import os
import sqlite3
import sys
import tempfile

# ---------------------------------------------------------------- 数据模型


class Block:
    __slots__ = ('num', 'density', 'sow')

    def __init__(self, num, density, sow):
        self.num, self.density, self.sow = num, density, sow


class MapDef:
    def __init__(self):
        self.name = ''
        self.rows = 0
        self.cols = 0
        self.block_col = 1     # QT BlockColumn 小区行数
        self.machine_cols = 6  # QT MachineColumn 机器行数
        self.col_gap = 60      # QT SeedColumnGap 行间距cm
        self.rowlen = []       # 每排行长cm
        self.aisle = []        # rows+1 条过道cm
        self.ab = None         # {name,sn,se,en,ee,sdir,cdir,a_row,abdir}
        self.blocks = []       # [[Block]]

    def validate(self):
        errs = []
        if not (1 <= self.rows <= 200 and 1 <= self.cols <= 200):
            errs.append('排/列数越界(%d×%d)' % (self.rows, self.cols))
        if not (1 <= self.machine_cols <= 16):
            errs.append('机器行数越界(%d)' % self.machine_cols)
        if not (1 <= self.block_col <= 8):
            errs.append('小区行数越界(%d)' % self.block_col)
        if self.machine_cols % self.block_col != 0:
            errs.append('机器行数%d不能被小区行数%d整除'
                        % (self.machine_cols, self.block_col))
        if len(self.rowlen) != self.rows:
            errs.append('#ROWLEN需%d条实有%d' % (self.rows, len(self.rowlen)))
        if len(self.aisle) != self.rows + 1:
            errs.append('#AISLE需%d条实有%d' % (self.rows + 1, len(self.aisle)))
        if self.ab:
            ab = self.ab
            for k in ('sn', 'se', 'en', 'ee', 'sdir', 'cdir'):
                if not isinstance(ab[k], (int, float)):
                    errs.append('#AB字段%s非数值' % k)
        return errs


# ---------------------------------------------------------------- slug(与store.lua一致)

def _fnv1a(s):
    h = 0x811C9DC5
    for b in s:
        h = ((h ^ b) & 0xFFFFFFFF) * 0x01000193 & 0xFFFFFFFF
    return h


def slug(name):
    ascii_part = []
    for ch in name:
        if ch.isascii() and ch.isalnum():
            ascii_part.append(ch)
        if len(ascii_part) >= 16:
            break
    head = ''.join(ascii_part) or 'n'
    return '%s_%08X' % (head, _fnv1a(name.encode('gb2312', errors='replace')))


# ---------------------------------------------------------------- fmt_num(与csv.lua一致)

def fmt_num(v):
    v = float(v)
    if v == math.floor(v) and abs(v) < 1e15:
        return '%d' % int(v)
    s = '%.8f' % v
    if '.' in s:
        s = s.rstrip('0').rstrip('.')
    return s


# ---------------------------------------------------------------- 渲染(与csv.lua render_map一致)

def render_map_csv(m):
    L = ['#MAP,%s' % m.name]
    if (m.block_col or 1) > 1:
        L.append('#DIM,%d,%d,%d,%s,%d' % (m.rows, m.cols, m.machine_cols,
                                          fmt_num(m.col_gap), m.block_col))
    else:
        L.append('#DIM,%d,%d,%d,%s' % (m.rows, m.cols, m.machine_cols,
                                       fmt_num(m.col_gap)))
    for r in range(m.rows):
        L.append('#ROWLEN,%s' % fmt_num(m.rowlen[r]))
    for i in range(m.rows + 1):
        L.append('#AISLE,%s' % fmt_num(m.aisle[i]))
    if m.ab:
        ab = m.ab
        L.append('#AB,%s,%s,%s,%s,%s,%s,%s,%d,%d' % (
            ab['name'], fmt_num(ab['sn']), fmt_num(ab['se']), fmt_num(ab['en']),
            fmt_num(ab['ee']), fmt_num(ab['sdir']), fmt_num(ab['cdir']),
            ab['a_row'], ab['abdir']))
    L.append('#BLOCKS')
    warn_nonnum = []
    for r in range(m.rows):
        for c in range(m.cols):
            b = m.blocks[r][c]
            num = b.num
            try:
                num_i = int(float(num))
            except (TypeError, ValueError):
                num_i = r * m.cols + c + 1          # 非数字编号 -> 顺位替代
                warn_nonnum.append('(%d,%d)%s' % (r + 1, c + 1, num))
            L.append('%d,%d,%d,%d,%d' % (r + 1, c + 1, num_i, b.density,
                                         1 if b.sow else 0))
    return '\r\n'.join(L) + '\r\n', warn_nonnum


# ---------------------------------------------------------------- 读取: SQLite

def read_maps_from_db(db_path):
    con = sqlite3.connect(db_path)
    con.row_factory = sqlite3.Row
    try:
        maps = []
        for r in con.execute('SELECT * FROM Map ORDER BY Name'):
            m = MapDef()
            m.name = r['Name']
            m.block_col = r['BlockColumn'] or 1
            m.machine_cols = r['MachineColumn'] or 6
            m.col_gap = r['SeedColumnGap'] or 60
            m.rows = con.execute(
                'SELECT COUNT(*) FROM RowAttributes WHERE MapName=?',
                (m.name,)).fetchone()[0]
            m.cols = con.execute(
                'SELECT MAX([Column])+1 FROM MapBlock WHERE MapName=?',
                (m.name,)).fetchone()[0] or 0
            m.rowlen = [x[0] for x in con.execute(
                'SELECT Length FROM RowAttributes WHERE MapName=? ORDER BY [Row]',
                (m.name,))]
            m.aisle = [x[0] for x in con.execute(
                'SELECT Length FROM MapAisle WHERE MapName=? ORDER BY [Index]',
                (m.name,))]
            m.blocks = [[Block('', 0, False) for _ in range(m.cols)]
                        for _ in range(m.rows)]
            for br in con.execute(
                    'SELECT [Row],[Column],Number,Density,SowSeed FROM MapBlock '
                    'WHERE MapName=?', (m.name,)):
                rr, cc = br['Row'], br['Column']
                if 0 <= rr < m.rows and 0 <= cc < m.cols:
                    m.blocks[rr][cc] = Block(br['Number'], br['Density'] or 0,
                                             bool(br['SowSeed']))
            abname = r['ABLineName']
            if abname:
                abr = con.execute('SELECT * FROM ABLine WHERE Name=?',
                                  (abname,)).fetchone()
                if abr:
                    m.ab = {'name': abname,
                            'sn': abr['StartN'], 'se': abr['StartE'],
                            'en': abr['EndBN'], 'ee': abr['EndBE'],
                            'sdir': abr['StartDir'], 'cdir': abr['CalcDir'],
                            'a_row': r['ARow'] or 0,
                            'abdir': 1 if r['ABLineDirection'] else 0}
            maps.append(m)
        return maps
    finally:
        con.close()


# ---------------------------------------------------------------- 读取: xlsx

def read_map_from_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    base = os.path.basename(path)
    name = base[2:].rsplit('.', 1)[0] if base.startswith('地图') \
        else os.path.splitext(base)[0]

    def cell(ws, r, c):
        v = ws.cell(row=r, column=c).value
        return v

    ws = wb['信息']
    m = MapDef()
    m.name = name
    m.rows = int(cell(ws, 5, 2))
    m.cols = int(cell(ws, 6, 2))
    m.block_col = int(cell(ws, 7, 2) or 1)
    m.machine_cols = int(cell(ws, 8, 2) or 6)
    m.col_gap = float(cell(ws, 9, 2) or 60)
    m.rowlen = [float(cell(ws, i + 2, 3)) for i in range(m.rows)]
    m.aisle = [float(cell(ws, i + 2, 4)) for i in range(m.rows + 1)]

    if 'AB线' in wb.sheetnames:
        wsa = wb['AB线']
        vals = {}
        for r in range(1, 17):
            k = cell(wsa, r, 1)
            if k:
                vals[k] = cell(wsa, r, 2)
        m.ab = {'name': str(vals.get('名称', name)),
                'sn': vals.get('StartN', 0.0), 'se': vals.get('StartE', 0.0),
                'en': vals.get('EndBN', 0.0), 'ee': vals.get('EndBE', 0.0),
                'sdir': vals.get('StartDir', 0.0), 'cdir': vals.get('CalcDir', 0.0),
                'a_row': int(vals.get('A点排', 0) or 0),
                'abdir': 1 if vals.get('方向') else 0}

    def grid(sheet, conv):
        # QXlsx write(i+2,j+2) 为1基 -> 数据始于B2(头行/头列是标签),
        # openpyxl 同为1基: cell(i+2, j+2) (真实文件核对 2026-09)
        wsg = wb[sheet]
        out = [[None] * m.cols for _ in range(m.rows)]
        for i in range(m.rows):
            for j in range(m.cols):
                out[i][j] = conv(wsg.cell(row=i + 2, column=j + 2).value)
        return out

    nums = grid('编号', lambda v: '' if v is None else str(v))
    dens = grid('密度', lambda v: int(v or 0))
    sows = grid('是否播种', lambda v: v is not None and str(v).strip() != '')
    m.blocks = [[Block(nums[i][j], dens[i][j], sows[i][j])
                 for j in range(m.cols)] for i in range(m.rows)]
    return m


# ---------------------------------------------------------------- 导出

def export_one(m, out_dir):
    errs = m.validate()
    if errs:
        return None, '校验失败: ' + '; '.join(errs)
    try:
        text, warns = render_map_csv(m)
        text.encode('gb2312')               # 名称必须可GB2312编码
    except UnicodeEncodeError:
        return None, '地图名/编号含GB2312无法编码的字符'
    fname = 'map_%s.csv' % slug(m.name)
    path = os.path.join(out_dir, fname)
    with open(path, 'wb') as f:
        f.write(text.encode('gb2312'))
    return path, ('; '.join(warns) if warns else None)


# ---------------------------------------------------------------- CLI

def run_cli(args):
    maps = []
    if args.db:
        if not os.path.exists(args.db):
            print('数据库不存在: %s (QT上位机运行目录下的 data.db)' % args.db)
            return 1
        maps = read_maps_from_db(args.db)
    if args.xlsx:
        for p in args.xlsx:
            maps.append(read_map_from_xlsx(p))
    if args.map:
        want = set(args.map)
        maps = [m for m in maps if m.name in want]
    if not maps:
        print('未找到地图 (db=%s xlsx=%d个)' % (args.db, len(args.xlsx or [])))
        return 1
    os.makedirs(args.out, exist_ok=True)
    fail = 0
    for m in maps:
        path, msg = export_one(m, args.out)
        if path:
            print('[OK] %-16s -> %s%s' % (m.name, os.path.basename(path),
                                          ('  | ' + msg) if msg else ''))
        else:
            print('[失败] %-16s : %s' % (m.name, msg))
            fail += 1
    return 1 if fail else 0


# ---------------------------------------------------------------- GUI

def run_gui():
    from PySide6.QtWidgets import (QApplication, QMainWindow, QWidget, QVBoxLayout,
                                   QHBoxLayout, QGridLayout, QPushButton, QLabel,
                                   QLineEdit, QListWidget, QListWidgetItem,
                                   QPlainTextEdit, QFileDialog, QMessageBox,
                                   QCheckBox, QGroupBox)
    from PySide6.QtCore import Qt

    class Win(QMainWindow):
        def __init__(self):
            super().__init__()
            self.setWindowTitle('QT地图 → 组态屏CSV 导出工具')
            self.maps = []
            central = QWidget()
            self.setCentralWidget(central)
            v = QVBoxLayout(central)

            g1 = QGridLayout()
            self.ed_db = QLineEdit()
            b_db = QPushButton('选择data.db…')
            b_db.clicked.connect(self.pick_db)
            self.ed_xlsx = QLineEdit()
            self.ed_xlsx.setPlaceholderText('或选择 地图*.xlsx (可多选, 分号分隔)')
            b_xlsx = QPushButton('选择xlsx…')
            b_xlsx.clicked.connect(self.pick_xlsx)
            g1.addWidget(QLabel('QT数据库:'), 0, 0)
            g1.addWidget(self.ed_db, 0, 1)
            g1.addWidget(b_db, 0, 2)
            g1.addWidget(QLabel('QT的xlsx:'), 1, 0)
            g1.addWidget(self.ed_xlsx, 1, 1)
            g1.addWidget(b_xlsx, 1, 2)
            v.addLayout(g1)

            b_load = QPushButton('读取地图列表')
            b_load.clicked.connect(self.load)
            v.addWidget(b_load)

            self.lst = QListWidget()
            v.addWidget(self.lst)

            g2 = QHBoxLayout()
            self.ed_out = QLineEdit(os.path.join(os.getcwd(), 'hmi_out'))
            b_out = QPushButton('输出目录…')
            b_out.clicked.connect(self.pick_out)
            b_all = QPushButton('全选/全不选')
            b_all.clicked.connect(self.toggle_all)
            g2.addWidget(QLabel('输出到:'))
            g2.addWidget(self.ed_out, 1)
            g2.addWidget(b_out)
            g2.addWidget(b_all)
            v.addLayout(g2)

            g3 = QHBoxLayout()
            b_export = QPushButton('导出所选')
            b_export.clicked.connect(self.export_sel)
            b_preview = QPushButton('预览选中CSV')
            b_preview.clicked.connect(self.preview)
            g3.addWidget(b_export)
            g3.addWidget(b_preview)
            v.addLayout(g3)

            self.txt = QPlainTextEdit()
            self.txt.setReadOnly(True)
            v.addWidget(self.txt, 1)
            self.statusBar().showMessage(
                '导出的CSV复制到U盘即可导入屏端; 文件名自动ASCII化, 显示名保留中文')

        def pick_db(self):
            p, _ = QFileDialog.getOpenFileName(self, '选择QT上位机 data.db',
                                               '', 'SQLite (*.db)')
            if p:
                self.ed_db.setText(p)
                self.ed_xlsx.clear()

        def pick_xlsx(self):
            ps, _ = QFileDialog.getOpenFileNames(self, '选择QT导出的地图xlsx',
                                                 '', 'Excel (*.xlsx)')
            if ps:
                self.ed_xlsx.setText(';'.join(ps))
                self.ed_db.clear()

        def pick_out(self):
            p = QFileDialog.getExistingDirectory(self, '输出目录')
            if p:
                self.ed_out.setText(p)

        def toggle_all(self):
            st = not (self.lst.count() and self.lst.item(0).checkState() == Qt.Checked)
            for i in range(self.lst.count()):
                self.lst.item(i).setCheckState(
                    Qt.Checked if st else Qt.Unchecked)

        def load(self):
            self.maps = []
            try:
                if self.ed_db.text().strip():
                    self.maps = read_maps_from_db(self.ed_db.text().strip())
                elif self.ed_xlsx.text().strip():
                    for p in self.ed_xlsx.text().split(';'):
                        p = p.strip()
                        if p:
                            self.maps.append(read_map_from_xlsx(p))
                else:
                    QMessageBox.information(self, '提示', '先选择 data.db 或 xlsx')
                    return
            except Exception as exc:
                QMessageBox.critical(self, '读取失败', str(exc))
                return
            self.lst.clear()
            for m in self.maps:
                errs = m.validate()
                it = QListWidgetItem('%s  %d排×%d列%s%s' % (
                    m.name, m.rows, m.cols,
                    ('  [AB:%s]' % m.ab['name']) if m.ab else '',
                    ('  [校验:%s]' % ';'.join(errs)) if errs else ''))
                it.setCheckState(Qt.Checked if not errs else Qt.Unchecked)
                self.lst.addItem(it)
            self.statusBar().showMessage('读取到 %d 张地图' % len(self.maps))

        def selected(self):
            out = []
            for i in range(self.lst.count()):
                if self.lst.item(i).checkState() == Qt.Checked:
                    out.append(self.maps[i])
            return out

        def preview(self):
            sel = self.selected()
            if not sel:
                return
            text, _ = render_map_csv(sel[0])
            self.txt.setPlainText(text.replace('\r\n', '\n'))

        def export_sel(self):
            sel = self.selected()
            if not sel:
                QMessageBox.information(self, '提示', '未选择地图')
                return
            out = self.ed_out.text().strip()
            os.makedirs(out, exist_ok=True)
            lines = []
            for m in sel:
                path, msg = export_one(m, out)
                if path:
                    lines.append('[OK] %s -> %s%s' % (
                        m.name, os.path.basename(path),
                        ('  | ' + msg) if msg else ''))
                else:
                    lines.append('[失败] %s : %s' % (m.name, msg))
            self.txt.setPlainText('\n'.join(lines))
            self.statusBar().showMessage('导出完成: %d/%d 成功'
                                         % (len(sel) - sum('失败' in x for x in lines),
                                            len(sel)))

    app = QApplication(sys.argv)
    w = Win()
    w.resize(760, 560)
    w.show()
    sys.exit(app.exec())


# ---------------------------------------------------------------- 自测

def selftest():
    """临时SQLite+自建xlsx -> 导出 -> 用HIM工程csv.lua真实解析器回读验证。"""
    print('1) 构造QT数据库(5排2列, 绑定AB线)')
    tmp = tempfile.mkdtemp()
    db = os.path.join(tmp, 'data.db')
    con = sqlite3.connect(db)
    con.executescript('''
    CREATE TABLE Map(Name TEXT PRIMARY KEY,CreateTime TEXT,ModifyTime TEXT,Mark TEXT,
        BlockColumn INT,MachineColumn INT,SeedColumnGap INT,
        ARow INT,AColumn INT,ABLineDirection INT,ABLineName TEXT);
    CREATE TABLE MapBlock(MapName TEXT,[Row] INT,[Column] INT,Number TEXT,
        Density INT,SowSeed INT);
    CREATE TABLE RowAttributes(MapName TEXT,[Row] INT,Length INT);
    CREATE TABLE MapAisle(MapName TEXT,[Index] INT,Length INT);
    CREATE TABLE ABLine(Name TEXT PRIMARY KEY,StartN REAL,StartE REAL,EndBN REAL,
        EndBE REAL,StartDir REAL,CalcDir REAL,StartState INT,CalcDirFlag INT,
        OffLen INT,BDOffTime INT,Resv INT,CrcCheck INT);
    ''')
    con.execute("INSERT INTO Map VALUES('对照田','t','t','',1,6,60,0,0,1,'ab1')")
    con.execute("INSERT INTO ABLine VALUES('ab1',36.4,117.0,0,0,0,0,1,0,60,100,0,0)")
    for r in range(5):
        con.execute("INSERT INTO RowAttributes VALUES('对照田',?,500)", (r,))
    for i in range(6):
        con.execute("INSERT INTO MapAisle VALUES('对照田',?,100)", (i,))
    for r in range(5):
        for c in range(2):
            con.execute("INSERT INTO MapBlock VALUES('对照田',?,?,?,?,1)",
                        (r, c, str(r * 2 + c + 1), 4000))
    con.commit()
    con.close()

    maps = read_maps_from_db(db)
    assert len(maps) == 1, '读出地图数'
    m = maps[0]
    ok0 = (m.rows == 5 and m.cols == 2 and m.machine_cols == 6
           and m.block_col == 1 and len(m.aisle) == 6 and m.ab is not None
           and m.ab['a_row'] == 0 and m.ab['abdir'] == 1)
    print('  [%s] SQLite读取: 5×2/机器6/过道6条/AB绑定 a_row=0 dir=同向'
          % ('通过' if ok0 else '失败'))

    text, warns = render_map_csv(m)
    assert '#MAP,对照田' in text and '#DIM,5,2,6,60' in text
    print('  [通过] 渲染: #MAP/#DIM 头部正确%s' % (' 警告:' + warns if warns else ''))

    print('2) HIM工程 csv.lua 真实解析器回读')
    try:
        from lupa import lua53 as lupa_lua
    except ImportError:
        print('  [跳过] 未安装lupa')
        return 0
    lua = lupa_lua.LuaRuntime(unpack_returned_tuples=True)
    him = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                        '..', '..', 'HIMproject', 'Project'))
    lua.globals().__P = os.path.join(him, 'csv.lua').replace('\\', '/')
    csv = lua.execute('local f = assert(loadfile(__P)); return f()')
    pm, msgs = csv['parse_map'](text)
    ok2 = (pm is not None
           and pm['name'] == '对照田' and pm['rows'] == 5 and pm['cols'] == 2
           and pm['machine_cols'] == 6
           and pm['blocks'][4][1]['density'] == 4000
           and pm['blocks'][0][0]['sow'] == 1
           and pm['ab'] is not None and pm['ab']['a_row'] == 0
           and pm['ab']['abdir'] == 1)
    print('  [%s] parse_map回读: 名称/维度/密度/播种/AB 全部正确%s' % (
        '通过' if ok2 else '失败',
        '' if pm is not None else ' msgs=%s' % msgs))

    print('3) xlsx读取路径(openpyxl构造同布局文件)')
    try:
        import openpyxl
    except ImportError:
        print('  [跳过] 未安装openpyxl')
        return 0
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '信息'
    ws.cell(5, 2, 5).value = 5
    ws.cell(6, 2, 2).value = 2
    ws.cell(7, 2, 1).value = 1
    ws.cell(8, 2, 6).value = 6
    ws.cell(9, 2, 60).value = 60
    for i in range(5):
        ws.cell(i + 2, 3).value = 500
    for i in range(6):
        ws.cell(i + 2, 4).value = 100
    wsa = wb.create_sheet('AB线')
    for r, (k, v) in enumerate([('StartN', 36.4), ('StartE', 117.0),
                                ('EndBN', 0), ('EndBE', 0), ('StartDir', 0),
                                ('CalcDir', 0), ('StartState', 1),
                                ('CalcDirFlag', 0), ('OffLen', 60),
                                ('BDOffTime', 100), ('Resv', 0), ('CrcCheck', 0),
                                ('A点排', 0), ('A点列', 0), ('方向', 1),
                                ('名称', 'ab1')], start=1):
        wsa.cell(r, 1).value = k
        wsa.cell(r, 2).value = v
    wsn = wb.create_sheet('编号')
    wsd = wb.create_sheet('密度')
    wss = wb.create_sheet('是否播种')
    # 夹具按真实QT布局: 数据始于B2(1基), 且数值非均匀(任何错位都会被抓到)
    for i in range(5):
        for j in range(2):
            wsn.cell(i + 2, j + 2).value = str(i * 2 + j + 1)
            wsd.cell(i + 2, j + 2).value = 4000 + i * 10 + j
            wss.cell(i + 2, j + 2).value = '√' if (i + j) % 2 == 0 else None
    xp = os.path.join(tmp, '地图对照田.xlsx')
    wb.save(xp)
    mx = read_map_from_xlsx(xp)
    # db侧也用同样的非均匀值, 两路径应逐字节一致
    m.blocks = [[Block(str(r * 2 + c + 1), 4000 + r * 10 + c, (r + c) % 2 == 0)
                 for c in range(2)] for r in range(5)]
    text, _ = render_map_csv(m)
    text2, _ = render_map_csv(mx)
    ok3 = text2 == text and mx.ab['a_row'] == 0 and mx.ab['abdir'] == 1
    print('  [%s] xlsx路径与db路径产出逐字节一致(非均匀值防错位)' % ('通过' if ok3 else '失败'))

    print('4) 文件名ASCII化')
    path, msg = export_one(m, tmp)
    fn = os.path.basename(path) if path else ''
    ok4 = (path is not None and fn == 'map_%s.csv' % slug('对照田')
           and all(ord(ch) < 127 for ch in fn)
           and open(path, 'rb').read().decode('gb2312') == text)
    print('  [%s] %s (纯ASCII, GB2312内容)' % ('通过' if ok4 else '失败', fn))

    n_fail = sum(1 for x in (ok0, ok2, ok3, ok4) if not x)
    print('自测%s (%d项失败)' % ('全部通过' if not n_fail else '有失败', n_fail))
    return 1 if n_fail else 0


def main():
    ap = argparse.ArgumentParser(description='QT地图 → 组态屏CSV 导出工具')
    ap.add_argument('--db', help='QT上位机 data.db 路径')
    ap.add_argument('--xlsx', nargs='+', help='QT导出的 地图*.xlsx 文件')
    ap.add_argument('--map', nargs='+', help='只导出指定名字的地图')
    ap.add_argument('--out', default='hmi_out', help='输出目录')
    ap.add_argument('--selftest', action='store_true', help='运行自测')
    ap.add_argument('--nogui', action='store_true', help='无参数时也用CLI(需--db/--xlsx)')
    args = ap.parse_args()
    if args.selftest:
        return selftest()
    if args.db or args.xlsx:
        return run_cli(args)
    if not args.nogui:
        try:
            return run_gui()
        except ImportError:
            print('未安装PySide6, 请用CLI模式: --db/--xlsx')
            return 1
    ap.print_help()
    return 0


if __name__ == '__main__':
    sys.exit(main())
